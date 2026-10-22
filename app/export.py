"""Excel export of the job list."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY = '"$"#,##0;[Red]-"$"#,##0'
STATUS_TEXT = {"ok": "On track", "trending": "Trending over", "over": "Over budget"}
STATUS_FILL = {
    "ok": PatternFill("solid", fgColor="D9F2D9"),
    "trending": PatternFill("solid", fgColor="FDEFC8"),
    "over": PatternFill("solid", fgColor="F6D3D3"),
}


def jobs_workbook(jobs: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    headers = ["Job #", "Job", "Customer", "% Complete", "Status"]
    for label in ("Labor", "Material", "Subcontractor", "Direct/Misc"):
        headers += [f"{label} Est", f"{label} Actual", f"{label} Over/(Under)"]
    headers += ["Total Est Cost", "Total Actual Cost", "Contract", "Invoiced", "Est Profit", "Projected Profit", "Last Sync"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for j in jobs:
        o = j["overall"]
        row = [j["job_number"], j["name"], j["customer"], j["pct_complete"], STATUS_TEXT[j["status"]]]
        for b in j["buckets"]:
            row += [b["estimate"], b["actual"], b["variance"]]
        row += [o["estimate"], o["actual"], o["contract"], o["invoiced"], o["estimated_profit"], o["projected_profit"], j["synced_at"]]
        ws.append(row)
        ws.cell(ws.max_row, 5).fill = STATUS_FILL[j["status"]]

    for col in range(6, len(headers)):
        for cell in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            cell[0].number_format = MONEY
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 2)
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 22
    ws.freeze_panes = "D2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
